import argparse
import json
import os
import sys
from pathlib import Path

import pandas as pd

from evaluator.core import SUPPORTED, collect_submission, evaluate, load_problem
from evaluator.report_generator import write_reports

PROJECT_ALIASES = [
    (("retail_purchase", "retail_customer_purchase", "customer purchase intelligence", "user_product_purchase_details"), "Retail_Customer_Purchase"),
    (("healthcare", "liver_disease", "liver disease", "liver patient dataset"), "Healthcare_Liver_Disease"),
    (("forecasting", "sales_forecasting", "retail sales forecasting", "retail_data.csv"), "Retail_Sales_Forecasting"),
    (("loan_default", "loandefault", "loan default", "default prediction"), "Loan_Default"),
]

def choose_project_milestone(path, config):
    path_text = path.as_posix().lower().replace("-", "_")
    filename = path.stem.lower().replace("-", "_")
    project_configs = config.get("projects", {})
    project = None

    # Prefer the project name in the file/folder path.
    for keywords, configured_name in PROJECT_ALIASES:
        if configured_name in project_configs and any(k in filename or k in path_text for k in keywords):
            project = configured_name
            break

    # If not found, inspect notebook/source text for project-specific evidence.
    if project is None and path.suffix.lower() in {".ipynb", ".py", ".md", ".txt", ".json", ".html"}:
        try:
            sample = path.read_text(encoding="utf-8", errors="replace")[:30000].lower()
            for keywords, configured_name in PROJECT_ALIASES:
                if configured_name in project_configs and any(k in sample for k in keywords):
                    project = configured_name
                    break
        except OSError:
            pass

    # Existing legacy submissions under the participant structure belong to Loan Default
    # only when the notebook itself or parent path indicates that project.
    if project is None:
        raise ValueError(
            f"Cannot identify project for '{path}'. Put a project keyword/title in the filename "
            "or notebook content. The evaluator will not silently assign the wrong project."
        )

    milestone = None
    for i in range(1, 5):
        if (f"milestone_{i}" in path_text or f"milestone{i}" in path_text
                or f"_m{i}_" in filename or f"_m{i}." in filename):
            milestone = f"Milestone_{i}"
            break
    if milestone is None:
        raise ValueError(f"Cannot identify milestone for '{path}'. Expected Milestone_1.ipynb through Milestone_4.ipynb.")

    if milestone not in project_configs[project].get("rubrics", {}):
        raise ValueError(f"No rubric configured for {project}/{milestone} in evaluator/config.json")
    return project, milestone

def discover_submissions(files_from=None):
    """Scan all submission folders on every run; do not restrict evaluation to changed files."""
    roots = [
        Path("submission"),
        Path("Submission-Milestone1"),
        Path("Submission-Milestone2"),
        Path("Submission-Milestone3"),
        Path("Submission-Milestone4"),
    ]
    found = set()
    for root in roots:
        if root.exists():
            for p in root.rglob("*"):
                if p.is_file() and (p.suffix.lower() in SUPPORTED or p.suffix.lower() == ".zip"):
                    found.add(p)
    # Keep files-from as a fallback for repositories with a custom submission root.
    if not found and files_from and Path(files_from).exists():
        for line in Path(files_from).read_text(encoding="utf-8").splitlines():
            p = Path(line.strip())
            if p.exists() and p.is_file() and (p.suffix.lower() in SUPPORTED or p.suffix.lower() == ".zip"):
                found.add(p)
    return sorted(found, key=lambda p: p.as_posix().lower())

def participant_name(path):
    p = Path(path)
    if len(p.parts) >= 3 and p.parts[0].lower() == "submission":
        return p.parts[1]
    import re
    match = re.match(r"(Participant[_ -][A-Za-z0-9]+)", p.stem, re.IGNORECASE)
    return match.group(1).replace(" ", "_").replace("-", "_") if match else p.stem

def write_milestone_reports(results, output_dir):
    """Write one consolidated Excel workbook for each milestone."""
    out = Path(output_dir)
    for i in range(1, 5):
        milestone = f"Milestone_{i}"
        milestone_results = [r for r in results if r["Milestone"] == milestone]
        folder = out / f"Submission-Milestone{i}-Report"
        folder.mkdir(parents=True, exist_ok=True)
        workbook_path = folder / f"Submission-Milestone{i}-Report.xlsx"

        fixed = ["Participant Name", "Submission File", "Project", "Milestone"]
        feedback = ["Total Score", "Maximum Score", "Area of Strength", "Area for Improvement", "Overall Feedback"]
        rows = []
        criterion_columns = []
        for item in milestone_results:
            result = item["result"]
            row = {
                "Participant Name": item["participant"],
                "Submission File": item["source"],
                "Project": result.get("project_id", ""),
                "Milestone": result.get("milestone_id", ""),
            }
            for c in result.get("criteria", []):
                col = f'{c.get("name", "Criterion")} ({c.get("max_marks", 0)})'
                row[col] = c.get("marks_awarded", 0)
                if col not in criterion_columns:
                    criterion_columns.append(col)
            row.update({
                "Total Score": result.get("total_marks", ""),
                "Maximum Score": result.get("maximum_marks", ""),
                "Area of Strength": text_value(result.get("strengths", [])),
                "Area for Improvement": text_value(result.get("areas_for_improvement", [])),
                "Overall Feedback": result.get("overall_feedback", ""),
            })
            rows.append(row)

        columns = fixed + criterion_columns + feedback
        frame = pd.DataFrame(rows, columns=columns)
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            frame.to_excel(writer, index=False, sheet_name="Evaluation Report")
            ws = writer.book["Evaluation Report"]
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in ws.columns:
                letter = get_column_letter(col[0].column)
                max_len = max((len(str(c.value or "")) for c in col), default=12)
                ws.column_dimensions[letter].width = min(48, max(14, max_len + 2))
                for cell in col:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        print(f"Milestone report: {workbook_path} ({len(rows)} submissions)")

def text_value(value):
    if isinstance(value, (list, tuple)):
        return "\\n".join(str(x) for x in value)
    return str(value or "")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--files-from", required=False, default="changed_submission_files.txt")
    ap.add_argument("--output-dir", default="evaluation-reports")
    args = ap.parse_args()

    config = json.loads(Path("evaluator/config.json").read_text(encoding="utf-8"))
    candidates = discover_submissions(args.files_from)
    if not candidates:
        print("No supported submission files found in submission folders.")
        return 0
    if not os.getenv("OPENAI_API_KEY"):
        print("ERROR: OPENAI_API_KEY secret is missing.", file=sys.stderr)
        return 2

    results = []
    failures = 0
    for path in candidates:
        try:
            project, milestone = choose_project_milestone(path, config)
            project_config = config["projects"][project]
            rubric_path = project_config["rubrics"][milestone]
            if not Path(rubric_path).exists():
                raise FileNotFoundError(f"Rubric file not found: {rubric_path}")
            rubric = json.loads(Path(rubric_path).read_text(encoding="utf-8"))
            if not rubric.get("criteria"):
                raise ValueError(f"Rubric has no criteria: {rubric_path}")
            problem_path = project_config.get("problem_statement")
            if not problem_path or not Path(problem_path).exists():
                raise FileNotFoundError(f"Problem statement file not found: {problem_path}")
            problem = load_problem(problem_path)
            evidence, included = collect_submission(
                [path],
                int(config.get("max_file_chars", 18000)),
                int(config.get("max_total_chars", 45000)),
            )
            if not evidence.strip():
                print(f"Skipping {path}: no readable submission content.")
                continue
            result = evaluate(evidence, rubric, problem, config.get("model", "gpt-4.1-mini"))
            # Keep project/milestone metadata authoritative from the selected configuration.
            result["project_id"] = project
            result["milestone_id"] = milestone
            write_reports(result, path.as_posix(), args.output_dir)
            results.append({
                "participant": participant_name(path),
                "source": path.as_posix(),
                "result": result,
                "Milestone": milestone,
            })
            print(f"Evaluated {path}: {result.get('total_marks', 0)}/{result.get('maximum_marks', 100)}")
        except Exception as exc:
            failures += 1
            print(f"ERROR evaluating {path}: {type(exc).__name__}: {exc}", file=sys.stderr)

    write_milestone_reports(results, args.output_dir)
    print(f"Completed: {len(results)} evaluated, {failures} failed.")
    return 1 if failures else 0

if __name__ == "__main__":
    raise SystemExit(main())
