import json
import re
from pathlib import Path
import pandas as pd

def write_reports(result, source_file, output_dir):
    """Write individual JSON/XLSX reports and update a consolidated feedback register."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    source = Path(source_file)
    stem = source.stem
    payload = dict(result)
    payload["source_file"] = source_file

    js = out / f"{stem}_evaluation.json"
    js.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    criteria = result.get("criteria", [])
    criterion_rows = [{
        "Criterion": c.get("name", ""),
        "Maximum Marks": c.get("max_marks", 0),
        "Marks Awarded": c.get("marks_awarded", 0),
        "Evidence": c.get("evidence", ""),
        "Criterion Feedback": c.get("feedback", "")
    } for c in criteria]
    xlsx = out / f"{stem}_evaluation.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        pd.DataFrame(criterion_rows).to_excel(writer, index=False, sheet_name="Criteria")
        pd.DataFrame([
            ["Participant Name", _participant_name(source_file)],
            ["Project", result.get("project_id", "")],
            ["Milestone", result.get("milestone_id", "")],
            ["Total Marks", result.get("total_marks", "")],
            ["Maximum Marks", result.get("maximum_marks", "")],
            ["Area of Strength", _as_text(result.get("strengths", []))],
            ["Area for Improvement", _as_text(result.get("areas_for_improvement", []))],
            ["Overall Feedback", result.get("overall_feedback", "")],
            ["Source File", source_file],
            ["Human Review Required", "Yes"]
        ], columns=["Field", "Value"]).to_excel(writer, index=False, sheet_name="Summary")
        _format_workbook(writer.book)

    # One row per submission; each rubric criterion gets its own score column.
    row = {
        "Participant Name": _participant_name(source_file),
        "Submission File": source_file,
        "Project": result.get("project_id", ""),
        "Milestone": result.get("milestone_id", "")
    }
    for c in criteria:
        row[f'{c.get("name", "Criterion")} ({c.get("max_marks", 0)})'] = c.get("marks_awarded", 0)
    row.update({
        "Total Score": result.get("total_marks", ""),
        "Maximum Score": result.get("maximum_marks", ""),
        "Area of Strength": _as_text(result.get("strengths", [])),
        "Area for Improvement": _as_text(result.get("areas_for_improvement", [])),
        "Overall Feedback": result.get("overall_feedback", "")
    })
    register = out / "Participant_Feedback_Register.xlsx"
    _upsert_feedback_row(register, row)
    return xlsx, js

def _participant_name(source_file):
    p = Path(source_file)
    if len(p.parts) >= 3 and p.parts[0].lower() == 'submission':
        return p.parts[1]
    stem = p.stem
    match = re.match(r'(Participant[_ -][A-Za-z0-9]+(?:[_ -][A-Za-z0-9]+)?)', stem, re.IGNORECASE)
    if match:
        return match.group(1).replace(' ', '_').replace('-', '_')
    return stem

def _as_text(value):
    if isinstance(value, (list, tuple)):
        return "\n".join(str(x) for x in value)
    return str(value or "")

def _upsert_feedback_row(path, row):
    # If the same file is evaluated again, replace its row instead of duplicating it.
    if path.exists():
        old = pd.read_excel(path, sheet_name="Feedback Register").fillna("")
        records = old.to_dict(orient="records")
    else:
        records = []
    records = [r for r in records if str(r.get("Submission File", "")) != str(row["Submission File"])]
    records.append(row)
    frame = pd.DataFrame(records)
    # Keep identity columns first, then criterion columns, then total and feedback columns.
    fixed = ["Participant Name", "Submission File", "Project", "Milestone"]
    feedback = ["Total Score", "Maximum Score", "Area of Strength", "Area for Improvement", "Overall Feedback"]
    criterion_cols = [c for c in frame.columns if c not in fixed + feedback]
    frame = frame.reindex(columns=fixed + criterion_cols + feedback)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Feedback Register")
        _format_workbook(writer.book)

def _format_workbook(workbook):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max((len(str(c.value or "")) for c in col), default=12)
            ws.column_dimensions[letter].width = min(48, max(14, max_len + 2))
            for cell in col:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
